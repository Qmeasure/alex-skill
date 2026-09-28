#!/usr/bin/env python3
"""把联动 Excel 模型里面向报告的两张 sheet 导出成 tex。

    python scripts/export_excel_to_tex.py 模型.xlsx 报告目录/

读取：
  报告数字  A 宏名 | B 数值 | C 显示格式(可空) | D 前缀 | E 后缀 | F 说明
  报告表格  A 列是行类型：
            表  B 表ID  C 图表标题(空=不编号)  D 列宽 DXA，逗号分隔(可空)  E 模式 标准/紧凑/半宽
            头  B 起为表头
            行  B 起为数据行
            亮  同「行」，目标公司行，浅蓝底加粗
            注  B 为口径小字，跨全部列
            空行分隔两张表
写出：
  报告目录/data/keyfigures.tex   \\newcommand 宏
  报告目录/tables/<表ID>.tex     每张表一个文件

Excel 先用内嵌 xlsx skill 的 recalc.py 重算。读到公式没有缓存值、错误值或
重复宏名/表ID 时直接失败，不输出半成品。
"""

import argparse
import datetime as dt
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from zfj_tex import render_macros, render_table, macro_name_ok  # noqa: E402

NUM_SHEET = "报告数字"
TAB_SHEET = "报告表格"
ERROR_VALUES = {"#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!",
                "#SPILL!", "#CALC!", "#BUSY!", "#GETTING_DATA", "#FIELD!", "#UNKNOWN!"}
RESERVED = {"ReportType", "ReportDate", "CoverTitle", "CoverSubtitle", "CoverStock"}   # front/meta.tex 已用


class ExportError(Exception):
    pass


def _format_section(value, section):
    """按一个 Excel 格式段输出（value 已按段取绝对值或原值）。支持引号字面量前后缀。"""
    lit = re.fullmatch(r'("([^"]*)")?(.*?)("([^"]*)")?', section)
    prefix, core, suffix = lit.group(2) or "", lit.group(3), lit.group(5) or ""
    if core.startswith("(") and core.endswith(")"):
        return prefix + "(" + _format_section(value, core[1:-1]) + ")" + suffix
    if core in ("", "@", "General"):
        return prefix + suffix
    sign = core.startswith("+")
    core = core.lstrip("+")
    pct = core.endswith("%")
    if pct:
        core = core[:-1]
        value = value * 100
    m = re.fullmatch(r'(#,##)?0(?:\.(0+))?(x|X|倍)?', core)
    if not m:
        raise ExportError(f"不支持的数字格式 {section!r}，改用 0 / 0.0 / #,##0.00 / 0.0% 这类格式")
    decimals = len(m.group(2) or "")
    text = f"{value:,.{decimals}f}" if m.group(1) else f"{value:.{decimals}f}"
    if sign and value > 0:
        text = "+" + text
    return prefix + text + ("%" if pct else "") + (m.group(3) or "") + suffix


def format_value(value, fmt):
    """按 Excel 数字格式把值转成文本。支持「正;负;零」三段、引号字面量和研报常用格式。"""
    if value is None:
        return ""
    if isinstance(value, str):
        if value.strip() in ERROR_VALUES:
            raise ExportError(f"错误值 {value}")
        return value
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (dt.date, dt.datetime)):
        return value.strftime("%Y-%m-%d")
    fmt = fmt or "General"
    if fmt in ("General", "@"):
        if isinstance(value, float) and not value.is_integer():
            if len(repr(value).split(".")[1]) > 4:
                raise ExportError(f"数值 {value} 没有设显示格式，给单元格设 0.00 / 0.0% 这类格式")
            return repr(value)
        return str(int(value)) if isinstance(value, float) else str(value)
    sections = fmt.split(";")
    if value < 0 and len(sections) >= 2:
        return _format_section(-value, sections[1])
    if value == 0 and len(sections) >= 3:
        return _format_section(0, sections[2])
    return _format_section(value, sections[0])


def cell_text(vcell, fcell, fmt=None):
    if isinstance(fcell.value, str) and fcell.value.startswith("=") and vcell.value is None:
        raise ExportError(f"{fcell.coordinate} 是公式但没有缓存值，先跑 recalc.py")
    try:
        return format_value(vcell.value, fmt or vcell.number_format)
    except ExportError as exc:
        raise ExportError(f"{fcell.coordinate}: {exc}") from None


def read_numbers(ws_v, ws_f):
    pairs, seen = [], set()
    for row_v, row_f in zip(ws_v.iter_rows(min_row=2), ws_f.iter_rows(min_row=2)):
        name = row_v[0].value
        if name is None or str(name).strip() == "":
            continue
        name = str(name).strip()
        if not macro_name_ok(name):
            raise ExportError(f"{NUM_SHEET}!{row_v[0].coordinate} 宏名只能用英文字母：{name}")
        if name in seen:
            raise ExportError(f"{NUM_SHEET} 宏名重复：{name}")
        seen.add(name)
        if name in RESERVED:
            raise ExportError(f"{NUM_SHEET} 宏名 {name} 与 front/meta.tex 重名，换一个")
        get = lambda i: row_v[i].value if len(row_v) > i else None  # noqa: E731
        fmt = get(2)
        text = cell_text(row_v[1], row_f[1], str(fmt) if fmt else None)
        if text == "":
            raise ExportError(f"{NUM_SHEET}!B{row_v[1].row} 宏 {name} 的值是空的")
        pairs.append((name, f"{get(3) or ''}{text}{get(4) or ''}"))
    return pairs


def read_tables(ws_v, ws_f):
    tables, cur, seen = [], None, set()

    def close():
        if cur is not None:
            if not cur["head"] and cur["mode"] in ("标准", "standard"):
                raise ExportError(f"表 {cur['id']} 是正文表，缺少「头」行")
            tables.append(cur)

    for row_v, row_f in zip(ws_v.iter_rows(), ws_f.iter_rows()):
        kind = row_v[0].value
        kind = str(kind).strip() if kind is not None else ""
        vals = row_v[1:]
        fvals = row_f[1:]
        if kind == "":
            if any(c.value not in (None, "") for c in vals):
                raise ExportError(f"{TAB_SHEET}!A{row_v[0].row} 缺少行类型")
            close()
            cur = None
            continue
        if kind == "表":
            close()
            tid = str(vals[0].value or "").strip()
            if not re.fullmatch(r"[A-Za-z0-9_-]+", tid):
                raise ExportError(f"{TAB_SHEET}!B{row_v[0].row} 表ID只能用英文、数字、下划线：{tid!r}")
            if tid in seen:
                raise ExportError(f"表ID重复：{tid}")
            seen.add(tid)
            widths = str(vals[2].value or "").strip() if len(vals) > 2 else ""
            cur = {"id": tid, "caption": (vals[1].value or "") if len(vals) > 1 else "",
                   "widths": [int(float(x)) for x in widths.split(",") if x.strip()] or None,
                   "mode": (vals[3].value or "标准") if len(vals) > 3 else "标准",
                   "head": [], "rows": [], "hi": [], "note": None}
            continue
        if cur is None:
            raise ExportError(f"{TAB_SHEET}!A{row_v[0].row}「{kind}」行前面没有「表」行")
        texts = [cell_text(v, f) for v, f in zip(vals, fvals)]
        while texts and texts[-1] == "":
            texts.pop()
        if kind == "头":
            cur["head"] = texts
        elif kind in ("行", "亮"):
            ncols = len(cur["head"]) or len(texts)
            if len(texts) > ncols:
                raise ExportError(f"{TAB_SHEET}!A{row_v[0].row} 这一行有 {len(texts)} 个值，表头只有 {ncols} 列")
            texts = texts + [""] * (ncols - len(texts))
            if kind == "亮":
                cur["hi"].append(len(cur["rows"]))
            cur["rows"].append(texts)
        elif kind == "注":
            cur["note"] = texts[0] if texts else ""
        else:
            raise ExportError(f"{TAB_SHEET}!A{row_v[0].row} 未知行类型「{kind}」，只能是 表/头/行/亮/注")
    close()
    return tables


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("report_dir")
    args = ap.parse_args()
    xlsx = Path(args.xlsx)
    out = Path(args.report_dir)
    try:
        wb_v = load_workbook(xlsx, data_only=True)
        wb_f = load_workbook(xlsx, data_only=False)
        for name in (NUM_SHEET, TAB_SHEET):
            if name not in wb_v.sheetnames:
                raise ExportError(f"{xlsx.name} 缺少 sheet「{name}」")
        pairs = read_numbers(wb_v[NUM_SHEET], wb_f[NUM_SHEET])
        tables = read_tables(wb_v[TAB_SHEET], wb_f[TAB_SHEET])
        note = f"由 scripts/export_excel_to_tex.py 从 {xlsx.name} 生成"
        macros = render_macros(pairs, note)
        rendered = {t["id"]: render_table(t, f"{note}，表 {t['id']}") for t in tables}
    except (ExportError, ValueError) as exc:
        print(f"[FAIL] 导出失败：{exc}")
        return 1
    except Exception as exc:  # 文件损坏、锁文件等，给出可读原因而不是堆栈
        print(f"[FAIL] 导出失败：读取 {xlsx.name} 出错（{type(exc).__name__}: {exc}）")
        return 1
    (out / "data").mkdir(parents=True, exist_ok=True)
    (out / "tables").mkdir(parents=True, exist_ok=True)
    (out / "data" / "keyfigures.tex").write_text(macros, encoding="utf-8")
    for tid, text in rendered.items():
        (out / "tables" / f"{tid}.tex").write_text(text, encoding="utf-8")
    print(f"[PASS] 导出 {len(pairs)} 个宏、{len(tables)} 张表 → {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
